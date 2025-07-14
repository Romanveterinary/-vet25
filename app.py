import os
import datetime
import traceback
import uuid
import random
from functools import wraps
from urllib.parse import urlparse
from itertools import groupby
from collections import defaultdict
import json
import io
import zipfile
import requests

from werkzeug.utils import secure_filename
from flask import Flask, render_template, redirect, url_for, flash, request, send_file
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, extract, UniqueConstraint
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from PIL import Image, ImageDraw
import openpyxl
from openpyxl.styles import Font, Alignment
from markupsafe import Markup

# ====================================================================
# 1. Налаштування додатку Flask
# ====================================================================

basedir = os.path.abspath(os.path.dirname(__file__))
app = Flask(__name__,
            template_folder=os.path.join(basedir, 'templates'),
            static_folder=os.path.join(basedir, 'static'))

# Окремі папки для завантажень (для локальної розробки)
UPLOAD_FOLDER_PHOTOS = os.path.join(basedir, 'static', 'uploads', 'photos')
UPLOAD_FOLDER_DOCS = os.path.join(basedir, 'static', 'uploads', 'documents')
UPLOAD_FOLDER_SAFETY = os.path.join(basedir, 'static', 'uploads', 'safety_reports')
os.makedirs(UPLOAD_FOLDER_PHOTOS, exist_ok=True)
os.makedirs(UPLOAD_FOLDER_DOCS, exist_ok=True)
os.makedirs(UPLOAD_FOLDER_SAFETY, exist_ok=True)

ALLOWED_PHOTO_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}
ALLOWED_DOC_EXTENSIONS = {'doc', 'docx', 'xls', 'xlsx', 'pdf', 'txt'}

app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY') or 'дуже-секретний-ключ-для-розробки'
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL') or 'sqlite:///' + os.path.join(basedir, 'vet25.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False


# ====================================================================
# 2. Налаштування Google Cloud Storage з ДЕТАЛЬНОЮ ДІАГНОСТИКОЮ
# ====================================================================
print("================== ЗАПУСК КОНФІГУРАЦІЇ GCS ==================")
storage_client = None
bucket = None
GCS_AVAILABLE = False
try:
    from google.cloud import storage
    GCS_AVAILABLE = True
    print("Діагностика: бібліотека google-cloud-storage знайдена.")
except ImportError:
    print("Діагностика: ПОМИЛКА! Бібліотека google-cloud-storage не знайдена.")

BUCKET_NAME = os.environ.get('GCS_BUCKET_NAME')

if GCS_AVAILABLE and BUCKET_NAME:
    print(f"Діагностика: Знайдено змінну GCS_BUCKET_NAME: {BUCKET_NAME}")
    try:
        gcs_credentials_json_str = os.environ.get('GOOGLE_APPLICATION_CREDENTIALS_JSON')
        if gcs_credentials_json_str:
            print("Діагностика: Знайдено змінну GOOGLE_APPLICATION_CREDENTIALS_JSON.")
            gcs_credentials_json = json.loads(gcs_credentials_json_str)
            storage_client = storage.Client.from_service_account_info(gcs_credentials_json)
            print("Діагностика: Створено клієнт GCS з JSON змінної.")
        else:
            print("Діагностика: Змінна GOOGLE_APPLICATION_CREDENTIALS_JSON не знайдена, спроба стандартної аутентифікації.")
            storage_client = storage.Client()

        bucket = storage_client.get_bucket(BUCKET_NAME)
        print(f"!!! УСПІХ: Успішно підключено до Google Cloud Storage, бакет: {BUCKET_NAME} !!!")
    except Exception as e:
        print(f"!!! ПОМИЛКА ПІДКЛЮЧЕННЯ GCS: Не вдалося підключитися. {e} !!!")
        bucket = None # Переконуємось, що bucket = None у разі помилки
        traceback.print_exc()
else:
    print("!!! ПОПЕРЕДЖЕННЯ: GCS не налаштовано (немає GCS_BUCKET_NAME або бібліотеки). Використовується локальне сховище. !!!")
print("================== ЗАВЕРШЕННЯ КОНФІГУРАЦІЇ GCS ==================")


# ====================================================================
# 3. Ініціалізація та моделі даних
# ====================================================================

db = SQLAlchemy(app)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'auth.login'
login_manager.login_message = "Будь ласка, увійдіть, щоб отримати доступ до цієї сторінки."
login_manager.login_message_category = "info"

def allowed_file(filename, allowed_extensions):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in allowed_extensions

uk_months = {1:"Січень",2:"Лютий",3:"Березень",4:"Квітень",5:"Травень",6:"Червень",7:"Липень",8:"Серпень",9:"Вересень",10:"Жовтень",11:"Листопад",12:"Грудень"}
uk_weekdays = {0:"Понеділок",1:"Вівторок",2:"Середа",3:"Четвер",4:"П'ятниця",5:"Субота",6:"Неділя"}

def format_bilingual_date(date_obj):
    if not isinstance(date_obj,(datetime.date,datetime.datetime)):
        return date_obj
    uk_day_name = uk_weekdays.get(date_obj.weekday(),'')
    uk_month_name = uk_months.get(date_obj.month,'')
    uk_date_str = f"{uk_day_name}, {date_obj.day:02d} {uk_month_name} {date_obj.year}"
    en_date_str = date_obj.strftime('%A, %d %B %Y')
    return f"{uk_date_str} ({en_date_str})"

def nl2br(value):
    return Markup(value.replace('\n','<br>\n'))

app.jinja_env.filters['nl2br'] = nl2br
app.jinja_env.filters['bilingual_date'] = format_bilingual_date

class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    registration_date = db.Column(db.DateTime, default=lambda:datetime.datetime.now(datetime.UTC))
    role = db.Column(db.String(20), nullable=False, default='user')
    is_active = db.Column(db.Boolean, default=False, nullable=False)
    photos = db.relationship('Photo', backref='user', lazy='dynamic', cascade="all, delete-orphan")
    enterprises = db.relationship('Enterprise', backref='owner', lazy='dynamic', cascade="all, delete-orphan")
    comments = db.relationship('Comment', backref='author', lazy='dynamic', cascade="all, delete-orphan")
    reports = db.relationship('MonthlyReport', backref='author', lazy='dynamic', cascade="all, delete-orphan")
    documents = db.relationship('Document', backref='uploader', lazy='dynamic', cascade="all, delete-orphan")
    
    def set_password(self, password): self.password_hash = generate_password_hash(password)
    def check_password(self, password): return check_password_hash(self.password_hash, password)
    
    @property
    def is_admin(self): return self.role == 'admin'
    @property
    def is_manager(self): return self.role in ['manager','admin']

class Announcement(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    content = db.Column(db.Text, nullable=False)
    last_updated = db.Column(db.DateTime, default=datetime.datetime.now(datetime.UTC), onupdate=datetime.datetime.now(datetime.UTC))

class Enterprise(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(150), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    photos = db.relationship('Photo', backref='enterprise', lazy=True)

class Comment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    text = db.Column(db.Text, nullable=False)
    timestamp = db.Column(db.DateTime, index=True, default=lambda:datetime.datetime.now(datetime.UTC))
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    photo_id = db.Column(db.Integer, db.ForeignKey('photo.id'), nullable=True)
    photo = db.relationship('Photo', backref=db.backref('comments', lazy='dynamic', cascade="all, delete-orphan"))

class Photo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    enterprise_id = db.Column(db.Integer, db.ForeignKey('enterprise.id'), nullable=True)
    photo_type = db.Column(db.String(50), nullable=False, default='Не вказано')
    animal_species = db.Column(db.String(50), nullable=True)
    organ_type = db.Column(db.String(50), nullable=True)
    filename = db.Column(db.String(255), nullable=False)
    filepath = db.Column(db.String(512), nullable=False)
    upload_date = db.Column(db.DateTime, default=lambda:datetime.datetime.now(datetime.UTC))
    analyzed_filepath = db.Column(db.String(512), nullable=True)
    checked_for_trichinella = db.Column(db.Boolean, default=False)
    checked_for_anisakids = db.Column(db.Boolean, default=False)
    checked_for_cysticercosis = db.Column(db.Boolean, default=False)

class ReportData(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey('monthly_report.id'), nullable=False)
    animal_type = db.Column(db.String(50), nullable=False) 
    received = db.Column(db.Integer, default=0)
    diseases_registered = db.Column(db.Integer, default=0)
    disease_sybirka = db.Column(db.Integer, default=0)
    disease_tuberkuloz = db.Column(db.Integer, default=0)
    disease_brutseloz = db.Column(db.Integer, default=0)
    disease_lepto = db.Column(db.Integer, default=0)
    disease_beshykha_svynei = db.Column(db.Integer, default=0)
    disease_chuma_svynei = db.Column(db.Integer, default=0)
    disease_nezarazni = db.Column(db.Integer, default=0)
    disease_inshi_zarazni = db.Column(db.Integer, default=0)
    died_from_trauma = db.Column(db.Integer, default=0)
    sent_to_sanbiynia = db.Column(db.Integer, default=0)
    expert_sybirka = db.Column(db.Integer, default=0)
    expert_tuberkuloz_total = db.Column(db.Integer, default=0)
    expert_tuberkuloz_util = db.Column(db.Integer, default=0)
    expert_tuberkuloz_sanpererobka = db.Column(db.Integer, default=0)
    expert_leikoz = db.Column(db.Integer, default=0)
    expert_tsystytserkoz_finoz = db.Column(db.Integer, default=0)
    expert_ekhinokokoz = db.Column(db.Integer, default=0)
    expert_fastsioloz = db.Column(db.Integer, default=0)
    expert_brutseloz = db.Column(db.Integer, default=0)
    expert_trykhineloz = db.Column(db.Integer, default=0)
    expert_leptospiroz = db.Column(db.Integer, default=0)
    expert_inshi_zarazni = db.Column(db.Integer, default=0)
    expert_inshi_invaziyni = db.Column(db.Integer, default=0)
    expert_nezarazni_util = db.Column(db.Integer, default=0)
    expert_nezarazni_prompererobka = db.Column(db.Integer, default=0)
    expert_nezarazni_na_utylzavod = db.Column(db.Integer, default=0)
    expert_zneshkodzhennia_tush = db.Column(db.Integer, default=0)
    expert_utylzavod_holiv = db.Column(db.Integer, default=0)

class MonthlyReport(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    enterprise_id = db.Column(db.Integer, db.ForeignKey('enterprise.id'), nullable=False)
    report_year = db.Column(db.Integer, nullable=False)
    report_month = db.Column(db.Integer, nullable=False)
    submission_date = db.Column(db.DateTime, default=lambda: datetime.datetime.now(datetime.UTC))
    production_total = db.Column(db.Text, nullable=True)
    production_researched = db.Column(db.Text, nullable=True)
    waste_utilized_ton = db.Column(db.Float, default=0.0)
    shortcomings_violations = db.Column(db.Text, nullable=True)
    report_data = db.relationship('ReportData', backref='report', lazy='dynamic', cascade="all, delete-orphan")
    enterprise = db.relationship('Enterprise', backref=db.backref('reports', lazy=True))
    __table_args__ = (UniqueConstraint('enterprise_id', 'report_year', 'report_month', name='_enterprise_year_month_uc'),)

class Document(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text, nullable=True)
    filename = db.Column(db.String(255), nullable=False)
    filepath = db.Column(db.String(512), nullable=False)
    upload_date = db.Column(db.DateTime, default=lambda: datetime.datetime.now(datetime.UTC))
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    
class SafetyReport(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    unique_id = db.Column(db.String(36), unique=True, nullable=False, default=lambda: str(uuid.uuid4()))
    name = db.Column(db.String(150), nullable=False)
    phone = db.Column(db.String(20), nullable=False)
    location = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text, nullable=False)
    photo_filename = db.Column(db.String(255), nullable=False)
    photo_filepath = db.Column(db.String(512), nullable=False)
    submission_date = db.Column(db.DateTime, default=lambda: datetime.datetime.now(datetime.UTC))

# ====================================================================
# 4. Допоміжні функції для роботи з файлами
# ====================================================================

def save_file_photo(file_stream, content_type, original_filename, animal_species=None, organ_type=None):
    if bucket:
        try:
            filename = f"{uuid.uuid4()}-{secure_filename(original_filename)}"
            path_parts = ["user_uploads"]
            if animal_species: path_parts.append(secure_filename(animal_species))
            if organ_type: path_parts.append(secure_filename(organ_type))
            path_parts.append(filename)
            gcs_path = "/".join(path_parts)
            blob = bucket.blob(gcs_path)
            blob.upload_from_file(file_stream, content_type=content_type)
            blob.make_public()
            return blob.public_url
        except Exception as e:
            print(f"!!! ПОМИЛКА ЗБЕРЕЖЕННЯ ФАЙЛУ В GCS: {e} !!!")
            raise
    else:
        print("!!! ПОПЕРЕДЖЕННЯ: Використовується локальне збереження файлу! !!!")
        path_parts = ['uploads', 'photos']
        if animal_species: path_parts.append(secure_filename(animal_species))
        if organ_type: path_parts.append(secure_filename(organ_type))
        local_upload_path = os.path.join(app.static_folder, *path_parts)
        os.makedirs(local_upload_path, exist_ok=True)
        filename = secure_filename(original_filename)
        filepath = os.path.join(local_upload_path, filename)
        file_stream.seek(0)
        with open(filepath, 'wb') as f:
            f.write(file_stream.read())
        return os.path.join(*path_parts, filename).replace('\\', '/')

def delete_file_from_storage(filepath):
    if not filepath: return
    if filepath.startswith('http'):
        if not bucket: return
        try:
            parsed_url = urlparse(filepath)
            gcs_path = parsed_url.path.lstrip('/')
            blob_name = gcs_path.replace(f"{BUCKET_NAME}/", "", 1)
            blob = bucket.blob(blob_name)
            if blob.exists():
                blob.delete()
        except Exception as e:
            print(f"Помилка видалення файлу {filepath} з GCS: {e}")
    else:
        try:
            local_path = os.path.join(app.static_folder, filepath)
            if os.path.exists(local_path):
                os.remove(local_path)
        except Exception as e:
            print(f"Помилка видалення локального файлу {filepath}: {e}")

# ====================================================================
# 5. Декоратори, завантажувач користувача та глобальні функції
# ====================================================================
@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))

@app.context_processor
def inject_announcement():
    if current_user.is_authenticated:
        announcement = Announcement.query.order_by(Announcement.last_updated.desc()).first()
        return dict(announcement=announcement)
    return dict(announcement=None)

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            flash('Для доступу до цієї сторінки потрібні права адміністратора.', 'danger')
            return redirect(url_for('main.index'))
        return f(*args, **kwargs)
    return decorated_function

def manager_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_manager:
            flash('Для доступу до цієї сторінки потрібні права керівника або адміністратора.', 'danger')
            return redirect(url_for('main.index'))
        return f(*args, **kwargs)
    return decorated_function

# ====================================================================
# 6. Маршрути (Blueprints)
# ====================================================================
from flask import Blueprint

main_bp = Blueprint('main', __name__)
auth_bp = Blueprint('auth', __name__, url_prefix='/auth')
photo_bp = Blueprint('photo', __name__, url_prefix='/photo')
admin_bp = Blueprint('admin', __name__, url_prefix='/admin')
report_bp = Blueprint('report', __name__, url_prefix='/reports')
document_bp = Blueprint('document', __name__, url_prefix='/documents')
safety_bp = Blueprint('safety', __name__, url_prefix='/complaint')

# --- Маршрути для Safety (Скарги) ---
@safety_bp.route('/safety-check', methods=['GET', 'POST'])
def safety_check():
    if request.method == 'POST':
        name = request.form.get('name')
        phone = request.form.get('phone')
        location = request.form.get('location')
        description = request.form.get('description')
        file = request.files.get('file')

        errors = []
        if not name: errors.append("Поле \"Ваше ім'я\" є обов'язковим.")
        if not phone: errors.append("Поле \"Номер телефону\" є обов'язковим.")
        if not location: errors.append("Поле \"Місце придбання\" є обов'язковим.")
        if not description: errors.append("Поле \"Опишіть обставини\" є обов'язковим.")
        if not file or not file.filename: errors.append("Фото продукту є обов'язковим для завантаження.")

        if errors:
            for error in errors:
                flash(error, 'danger')
            return redirect(url_for('safety.safety_check'))

        try:
            filename = secure_filename(file.filename)
            save_path = os.path.join(UPLOAD_FOLDER_SAFETY, filename)
            file.save(save_path)
            relative_path = os.path.join('uploads', 'safety_reports', filename).replace('\\', '/')

            new_report = SafetyReport(
                name=name, phone=phone, location=location, description=description,
                photo_filename=filename, photo_filepath=relative_path
            )
            db.session.add(new_report)
            db.session.commit()
            
            flash('Ваше звернення успішно відправлено! Будь ласка, збережіть посилання для перевірки статусу.', 'success')
            return redirect(url_for('safety.submission_success', report_id=new_report.unique_id))
        except Exception as e:
            flash(f'Під час збереження форми сталася системна помилка: {e}', 'danger')
            traceback.print_exc()
            return redirect(url_for('safety.safety_check'))
    return render_template('safety_check.html')

@safety_bp.route('/submission-success/<string:report_id>')
def submission_success(report_id):
    inquiry_url = url_for('safety.inquiry_status', report_id=report_id, _external=True)
    return render_template('submission_success.html', inquiry_url=inquiry_url)

@safety_bp.route('/inquiry/<string:report_id>')
def inquiry_status(report_id):
    report = SafetyReport.query.filter_by(unique_id=report_id).first_or_404()
    return render_template('inquiry_status.html', report=report)

@safety_bp.route('/complaint-guide')
def complaint_guide():
    return render_template('complaint_guide.html')

# --- Основні маршрути ---
@main_bp.route('/')
def index():
    return render_template('index.html')

@main_bp.route('/my_photos')
@login_required
def my_photos():
    photos = Photo.query.filter_by(user_id=current_user.id).order_by(Photo.upload_date.desc()).all()
    return render_template('my_photos.html', photos=photos)

@main_bp.route('/community_feed')
@login_required
def community_feed():
    all_photos = Photo.query.order_by(Photo.upload_date.desc()).all()
    def get_date(photo):
        return photo.upload_date.date()
    grouped_photos = []
    for date, group in groupby(all_photos, key=get_date):
        grouped_photos.append((date, list(group)))
    return render_template('community_feed.html', grouped_photos=grouped_photos)

@main_bp.route('/my_enterprises', methods=['GET', 'POST'])
@login_required
def my_enterprises():
    if request.method == 'POST':
        name = request.form.get('enterprise_name')
        if name and not Enterprise.query.filter_by(name=name, user_id=current_user.id).first():
            db.session.add(Enterprise(name=name, user_id=current_user.id))
            db.session.commit()
            flash(f'Підприємство "{name}" успішно додано.', 'success')
        else:
            flash('Таке підприємство вже існує у вас або назва порожня.', 'danger')
        return redirect(url_for('main.my_enterprises'))
    enterprises = Enterprise.query.filter_by(user_id=current_user.id).order_by(Enterprise.name).all()
    return render_template('my_enterprises.html', enterprises=enterprises)

# --- Маршрути для автентифікації ---
@auth_bp.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('main.index'))
    if request.method == 'POST':
        user = User.query.filter_by(username=request.form['username']).first()
        if user and user.check_password(request.form['password']):
            if not user.is_active:
                flash('Ваш акаунт ще не активовано. Будь ласка, дочекайтеся схвалення адміністратором.', 'warning')
                return redirect(url_for('auth.login'))
            login_user(user)
            flash('Успішний вхід!', 'success')
            return redirect(url_for('main.index'))
        else:
            flash('Неправильне ім\'я користувача або пароль.', 'danger')
    return render_template('login.html')

@auth_bp.route('/register', methods=['GET', 'POST'])
def register():
    if current_user.is_authenticated:
        return redirect(url_for('main.index'))
    if request.method == 'POST':
        username = request.form.get('username')
        if User.query.filter_by(username=username).first():
            flash('Користувач з таким ім\'ям вже існує.', 'danger')
            return redirect(url_for('auth.register'))
        new_user = User(username=username, is_active=False)
        new_user.set_password(request.form.get('password'))
        db.session.add(new_user)
        db.session.commit()
        flash('Дякуємо за реєстрацію! Ваш акаунт буде активовано адміністратором найближчим часом.', 'info')
        return redirect(url_for('auth.login'))
    return render_template('register.html')

@auth_bp.route('/logout')
@login_required
def logout():
    logout_user()
    flash('Ви успішно вийшли.', 'info')
    return redirect(url_for('auth.login'))

# --- Маршрути для фото ---
@photo_bp.route('/upload', methods=['GET', 'POST'])
@login_required
def upload():
    enterprises = Enterprise.query.filter_by(user_id=current_user.id).order_by(Enterprise.name).all()
    if request.method == 'POST':
        file = request.files.get('file')
        enterprise_id = request.form.get('enterprise_id')
        photo_type = request.form.get('photo_type')
        animal_species = request.form.get('animal_species')
        organ_type = request.form.get('organ_type')
        if not all([file, file.filename, enterprise_id, photo_type]):
            flash('Будь ласка, заповніть усі поля та виберіть файл.', 'danger')
            return render_template('upload_photo.html', enterprises=enterprises)
        if allowed_file(file.filename, ALLOWED_PHOTO_EXTENSIONS):
            try:
                file.stream.seek(0)
                filepath = save_file_photo(file.stream, file.content_type, file.filename, animal_species, organ_type)
                new_photo = Photo(
                    user_id=current_user.id,
                    filename=secure_filename(file.filename),
                    filepath=filepath, 
                    enterprise_id=enterprise_id,
                    photo_type=photo_type,
                    animal_species=animal_species,
                    organ_type=organ_type,
                    checked_for_trichinella='check_trichinella' in request.form,
                    checked_for_anisakids='check_anisakids' in request.form,
                    checked_for_cysticercosis='check_cysticercosis' in request.form
                )
                db.session.add(new_photo)
                db.session.commit()
                flash(f'Файл {secure_filename(file.filename)} успішно завантажено!', 'success')
                return redirect(url_for('photo.upload'))
            except Exception as e:
                flash(f'Під час завантаження файлу сталася помилка: {e}', 'danger')
                return redirect(request.url)
        else:
            flash('Недопустимий тип файлу.', 'danger')
            return redirect(request.url)
    return render_template('upload_photo.html', enterprises=enterprises)

@photo_bp.route('/view/<int:photo_id>', methods=['GET', 'POST'])
@login_required
def view_details(photo_id):
    photo = db.session.get(Photo, photo_id)
    if not photo:
        flash('Фото не знайдено.', 'danger')
        return redirect(url_for('main.community_feed'))
    
    if request.method == 'POST':
        comment_text = request.form.get('comment_text')
        if comment_text:
            comment = Comment(text=comment_text, author=current_user, photo=photo)
            db.session.add(comment)
            db.session.commit()
            flash('Ваш коментар додано.', 'success')
            return redirect(url_for('photo.view_details', photo_id=photo_id))
        else:
            flash('Текст коментаря не може бути порожнім.', 'danger')
    
    comments = photo.comments.order_by(Comment.timestamp.asc()).all()
    return render_template('analysis_result.html', photo=photo, comments=comments)

@photo_bp.route('/comment/delete/<int:comment_id>', methods=['POST'])
@login_required
def delete_comment(comment_id):
    comment = db.session.get(Comment, comment_id)
    if not comment:
        flash('Коментар не знайдено.', 'danger')
        return redirect(url_for('main.community_feed'))

    photo_id = comment.photo_id
    if current_user.id == comment.user_id or current_user.is_admin:
        db.session.delete(comment)
        db.session.commit()
        flash('Коментар видалено.', 'success')
    else:
        flash('У вас немає прав для видалення цього коментаря.', 'danger')
    
    return redirect(url_for('photo.view_details', photo_id=photo_id))

@photo_bp.route('/delete/<int:photo_id>', methods=['POST'])
@login_required
@admin_required
def delete_photo(photo_id):
    photo = db.session.get(Photo, photo_id)
    if photo:
        delete_file_from_storage(photo.filepath)
        if photo.analyzed_filepath:
            delete_file_from_storage(photo.analyzed_filepath)
        
        db.session.delete(photo)
        db.session.commit()
        flash(f'Фото "{photo.filename}" та всі пов\'язані коментарі було видалено.', 'success')
        return redirect(url_for('main.community_feed'))
    else:
        flash('Фото не знайдено.', 'danger')
        return redirect(url_for('main.community_feed'))

@photo_bp.route('/analyze/<int:photo_id>')
@login_required
def perform_analysis(photo_id):
    photo = db.session.get(Photo, photo_id)
    if not photo or (photo.user_id != current_user.id and not current_user.is_admin):
        flash('Фото не знайдено або у вас немає доступу.', 'danger')
        return redirect(url_for('main.my_photos'))

    if photo.analyzed_filepath:
        return redirect(url_for('photo.view_details', photo_id=photo.id))

    try:
        print(f"Починаємо аналіз для фото ID: {photo_id}. Шлях до файлу: {photo.filepath}")
        if photo.filepath.startswith('http'):
            response = requests.get(photo.filepath, stream=True)
            response.raise_for_status() 
            image_stream = io.BytesIO(response.content)
            print("Фото успішно завантажено з URL.")
        else:
            print(f"ПОПЕРЕДЖЕННЯ: Аналіз локального файлу: {photo.filepath}. Це не повинно відбуватися на Render.")
            local_path = os.path.join(app.static_folder, photo.filepath)
            if not os.path.exists(local_path):
                flash(f'Помилка: вихідний файл не знайдено за шляхом {local_path}.', 'danger')
                return redirect(url_for('main.my_photos'))
            image_stream = open(local_path, 'rb')

        with Image.open(image_stream) as img:
            img_copy = img.convert("RGB")
            draw = ImageDraw.Draw(img_copy)
            width, height = img_copy.size
            for _ in range(random.randint(3, 5)):
                x, y = random.randint(0, width - 50), random.randint(0, height - 50)
                radius = random.randint(10, 25)
                draw.ellipse([x, y, x + radius*2, y + radius*2], outline='red', width=3)
            
            analyzed_buffer = io.BytesIO()
            img_copy.save(analyzed_buffer, format='JPEG')
            analyzed_buffer.seek(0)
            print("Емуляція аналізу (малювання кіл) завершена.")

        filename, ext = os.path.splitext(photo.filename)
        analyzed_filename = f"{filename}_analyzed.jpg"
        
        print("Спроба зберегти проаналізований файл...")
        analyzed_url = save_file_photo(
            analyzed_buffer, 'image/jpeg', analyzed_filename,
            photo.animal_species, photo.organ_type
        )
        print(f"Файл збережено, отримано URL: {analyzed_url}")
        
        photo.analyzed_filepath = analyzed_url
        db.session.commit()
        print("Шлях до проаналізованого файлу збережено в БД.")
        
        flash(f'Фото "{photo.filename}" успішно проаналізовано.', 'success')
        return redirect(url_for('photo.view_details', photo_id=photo.id))

    except Exception as e:
        print(f"!!! КРИТИЧНА ПОМИЛКА в perform_analysis: {e} !!!")
        traceback.print_exc()
        flash('Під час аналізу фото сталася критична помилка.', 'danger')
        return redirect(url_for('main.my_photos'))

@photo_bp.route('/download/<int:photo_id>')
@login_required
def download_file(photo_id):
    photo = db.session.get(Photo, photo_id)
    if not photo or (photo.user_id != current_user.id and not current_user.is_admin):
        flash('Фото не знайдено або у вас немає доступу.', 'danger')
        return redirect(request.referrer or url_for('main.index'))
    if photo.filepath.startswith('http'):
        return redirect(photo.filepath)
    try:
        full_path = os.path.join(basedir, 'static', photo.filepath)
        return send_file(full_path, as_attachment=True)
    except Exception as e:
        print(f"ПОМИЛКА при завантаженні файлу: {e}"); traceback.print_exc()
        flash('Не вдалося завантажити файл.', 'danger')
        return redirect(request.referrer)

# --- Маршрути для Адміна ---
@admin_bp.route('/panel')
@login_required
@manager_required
def panel():
    photos = Photo.query.order_by(Photo.upload_date.desc()).all()
    structured_photos = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    for photo in photos:
        pt = photo.photo_type or "Не вказано"
        animal = photo.animal_species or "Не вказано"
        organ = photo.organ_type or "Не вказано"
        structured_photos[pt][animal][organ].append(photo)
    return render_template('admin_panel.html', structured_photos=structured_photos)

@admin_bp.route('/announcement', methods=['GET', 'POST'])
@login_required
@manager_required
def manage_announcement():
    announcement = Announcement.query.first()
    if request.method == 'POST':
        content = request.form.get('content', '')
        if announcement:
            announcement.content = content
        else:
            announcement = Announcement(content=content)
            db.session.add(announcement)
        db.session.commit()
        flash('Оголошення успішно оновлено.', 'success')
        return redirect(url_for('admin.manage_announcement'))
    return render_template('admin_announcement.html', announcement=announcement)

@admin_bp.route('/complaints')
@login_required
@manager_required
def view_complaints():
    complaints = SafetyReport.query.order_by(SafetyReport.submission_date.desc()).all()
    return render_template('view_complaints.html', complaints=complaints)

@admin_bp.route('/reports')
@login_required
@manager_required
def reports():
    today = datetime.date.today()
    thirty_days_ago = today - datetime.timedelta(days=30)
    total_photos_week = db.session.query(func.count(Photo.id)).filter(Photo.upload_date >= (today - datetime.timedelta(days=7))).scalar()
    photos_today = db.session.query(func.count(Photo.id)).filter(func.date(Photo.upload_date) == today).scalar()
    most_common_pathology_query = db.session.query(Photo.animal_species, Photo.organ_type, func.count(Photo.id).label('count')).filter(Photo.upload_date >= thirty_days_ago, Photo.animal_species.isnot(None), Photo.organ_type.isnot(None)).group_by(Photo.animal_species, Photo.organ_type).order_by(func.count(Photo.id).desc()).first()
    most_common_pathology = f"{most_common_pathology_query[0]} - {most_common_pathology_query[1]}" if most_common_pathology_query else "Немає даних"
    uploads_per_day_query = db.session.query(func.date(Photo.upload_date).label('date'), func.count(Photo.id).label('count')).filter(Photo.upload_date >= thirty_days_ago).group_by(func.date(Photo.upload_date)).order_by(func.date(Photo.upload_date)).all()
    chart_data = {(today-datetime.timedelta(days=i)).strftime('%Y-%m-%d'): 0 for i in range(30)}
    for row in uploads_per_day_query:
        date_key = row.date.strftime('%Y-%m-%d') if isinstance(row.date, datetime.date) else row.date
        chart_data[date_key] = row.count
    chart_labels = json.dumps(list(chart_data.keys()))
    chart_values = json.dumps(list(chart_data.values()))
    users = User.query.filter(User.role != 'admin').all()
    user_stats = [{'username': user.username, 'photo_count': user.photos.filter(Photo.upload_date >= (today - datetime.timedelta(days=7))).count(), 'comment_count': user.comments.filter(Comment.timestamp >= (today - datetime.timedelta(days=7))).count()} for user in users]
    return render_template('reports.html', total_photos_week=total_photos_week, photos_today=photos_today, most_common_pathology=most_common_pathology, chart_labels=chart_labels, chart_values=chart_values, user_stats=user_stats)

@admin_bp.route('/download_report/<string:period>')
@login_required
@manager_required
def download_report(period):
    try:
        if period == 'week':
            days, period_name = 7, "тиждень"
        elif period == 'month':
            days, period_name = 30, "місяць"
        else:
            flash('Неправильний період для звіту.', 'danger')
            return redirect(url_for('admin.reports'))

        end_date = datetime.datetime.now(datetime.UTC)
        start_date = end_date - datetime.timedelta(days=days)

        # 1. Отримати дані
        photos = Photo.query.filter(Photo.upload_date.between(start_date, end_date)).all()
        comments = Comment.query.filter(Comment.timestamp.between(start_date, end_date)).all()

        # 2. Об'єднати та відсортувати події
        events = []
        for photo in photos:
            events.append({
                'timestamp': photo.upload_date,
                'user': photo.user.username,
                'action': 'Завантажив фото',
                'details': photo.filename
            })

        for comment in comments:
            events.append({
                'timestamp': comment.timestamp,
                'user': comment.author.username,
                'action': 'Написав коментар',
                'details': f'До фото ID {comment.photo_id}: "{comment.text[:100]}..."'
            })

        events.sort(key=lambda x: x['timestamp'])

        # 3. Створити Excel файл
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = f"Детальний звіт за {period_name}"

        header = ["Дата і час", "Користувач", "Дія", "Деталі"]
        sheet.append(header)
        header_font = Font(bold=True)
        for cell in sheet[1]:
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Додавання рядків з подіями
        for event in events:
            # Конвертація часу в локальний часовий пояс для зручності
            try:
                local_tz = datetime.datetime.now().astimezone().tzinfo
                local_time = event['timestamp'].astimezone(local_tz)
                formatted_time = local_time.strftime('%Y-%m-%d %H:%M:%S')
            except:
                formatted_time = event['timestamp'].strftime('%Y-%m-%d %H:%M:%S')

            row = [
                formatted_time,
                event['user'],
                event['action'],
                event['details']
            ]
            sheet.append(row)

        # Налаштування ширини колонок
        sheet.column_dimensions['A'].width = 20  # Дата і час
        sheet.column_dimensions['B'].width = 25  # Користувач
        sheet.column_dimensions['C'].width = 25  # Дія
        sheet.column_dimensions['D'].width = 80  # Деталі

        # Збереження у буфер
        virtual_workbook = io.BytesIO()
        workbook.save(virtual_workbook)
        virtual_workbook.seek(0)
        
        filename = f'detailed_report_{period}_{datetime.date.today()}.xlsx'
        return send_file(
            virtual_workbook,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        print(f"ПОМИЛКА при генерації Excel-звіту: {e}")
        traceback.print_exc()
        flash('Не вдалося згенерувати детальний звіт.', 'danger')
        return redirect(url_for('admin.reports'))


@admin_bp.route('/manage_enterprises', methods=['GET', 'POST'])
@login_required
@admin_required
def manage_enterprises():
    if request.method == 'POST':
        if not current_user.is_admin:
            flash('Тільки адміністратор може змінювати підприємства.', 'danger')
            return redirect(url_for('admin.manage_enterprises'))
        enterprise_id = request.form.get('enterprise_id')
        new_name = request.form.get('new_name')
        enterprise_to_update = db.session.get(Enterprise, enterprise_id)
        if enterprise_to_update and new_name:
            enterprise_to_update.name = new_name
            db.session.commit()
            flash(f'Назву підприємства оновлено на "{new_name}".', 'success')
        else:
            flash('Помилка оновлення назви.', 'danger')
        return redirect(url_for('admin.manage_enterprises'))
    all_enterprises = Enterprise.query.order_by(Enterprise.name).all()
    return render_template('admin_manage_enterprises.html', enterprises=all_enterprises)

@admin_bp.route('/users')
@login_required
@admin_required
def manage_users():
    users = User.query.order_by(User.is_active.desc(), User.registration_date.desc()).all()
    return render_template('admin_manage_users.html', users=users)

@admin_bp.route('/users/activate/<int:user_id>', methods=['POST'])
@login_required
@admin_required
def activate_user(user_id):
    user = db.session.get(User, user_id)
    if user:
        user.is_active = not user.is_active
        db.session.commit()
        status = "активовано" if user.is_active else "деактивовано"
        flash(f'Акаунт користувача {user.username} було {status}.', 'success')
    return redirect(url_for('admin.manage_users'))

@admin_bp.route('/users/reset_password/<int:user_id>', methods=['POST'])
@login_required
@admin_required
def admin_reset_password(user_id):
    user = db.session.get(User, user_id)
    if user:
        temp_password = ''.join(random.choices('abcdefghjkmnpqrstuvwxyz23456789', k=8))
        user.set_password(temp_password)
        db.session.commit()
        flash(f'Пароль для {user.username} скинуто. Новий пароль: {temp_password}', 'success')
    return redirect(url_for('admin.manage_users'))

@admin_bp.route('/users/delete/<int:user_id>', methods=['POST'])
@login_required
@admin_required
def admin_delete_user(user_id):
    user = db.session.get(User, user_id)
    if user and not user.is_admin:
        for photo in user.photos:
            delete_file_from_storage(photo.filepath)
            if photo.analyzed_filepath:
                delete_file_from_storage(photo.analyzed_filepath)
        db.session.delete(user)
        db.session.commit()
        flash(f'Користувача {user.username} та всі його дані було видалено.', 'success')
    else:
        flash('Неможливо видалити цього користувача.', 'danger')
    return redirect(url_for('admin.manage_users'))

@admin_bp.route('/enterprises/delete/<int:enterprise_id>', methods=['POST'])
@login_required
@admin_required
def delete_enterprise(enterprise_id):
    enterprise = db.session.get(Enterprise, enterprise_id)
    if enterprise:
        if enterprise.photos: 
            flash(f'Неможливо видалити підприємство "{enterprise.name}", оскільки до нього прив\'язані фотографії.', 'warning')
            return redirect(url_for('admin.manage_enterprises'))
        
        try:
            db.session.delete(enterprise)
            db.session.commit()
            flash(f'Підприємство "{enterprise.name}" успішно видалено.', 'success')
        except Exception as e:
            db.session.rollback()
            flash(f'Під час видалення підприємства сталася помилка: {e}', 'danger')
    else:
        flash('Підприємство не знайдено.', 'danger')
    return redirect(url_for('admin.manage_enterprises'))

@admin_bp.route('/archive/create')
@login_required
@manager_required
def create_archive():
    try:
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'a', zipfile.ZIP_DEFLATED, False) as zip_file:
            
            # --- Лог фотографій ---
            photos = Photo.query.order_by(Photo.upload_date.asc()).all()
            photos_log_data = []
            for photo in photos:
                photos_log_data.append({
                    'photo_id': photo.id,
                    'uploader_username': photo.user.username,
                    'enterprise_name': photo.enterprise.name if photo.enterprise else "Не вказано",
                    'filename': photo.filename,
                    'upload_date_utc': photo.upload_date.isoformat(),
                    'photo_type': photo.photo_type,
                    'animal_species': photo.animal_species,
                    'organ_type': photo.organ_type
                })
                
                if not photo.filepath.startswith('http'):
                    try:
                        full_path = os.path.join(app.static_folder, photo.filepath)
                        if os.path.exists(full_path):
                            arcname = os.path.join('photos', photo.user.username, str(photo.upload_date.year), str(photo.upload_date.month), photo.filename)
                            zip_file.write(full_path, arcname=arcname)
                    except Exception as e:
                        print(f"Не вдалося додати фото {photo.id} в архів: {e}")
            
            json_photos_log = json.dumps(photos_log_data, indent=4, ensure_ascii=False)
            zip_file.writestr('photos_log.json', json_photos_log)

            # --- Лог коментарів ---
            comments = Comment.query.order_by(Comment.timestamp.asc()).all()
            comments_data = []
            for comment in comments:
                comments_data.append({
                    'comment_id': comment.id,
                    'photo_id': comment.photo_id,
                    'author_username': comment.author.username,
                    'text': comment.text,
                    'timestamp_utc': comment.timestamp.isoformat()
                })
            
            json_comments_log = json.dumps(comments_data, indent=4, ensure_ascii=False)
            zip_file.writestr('comments_log.json', json_comments_log)

        zip_buffer.seek(0)
        archive_filename = f'archive_{datetime.date.today()}.zip'
        return send_file(zip_buffer, mimetype='application/zip', as_attachment=True, download_name=archive_filename)
        
    except Exception as e:
        flash(f'Під час створення архіву сталася помилка: {e}', 'danger')
        traceback.print_exc()
        return redirect(url_for('admin.panel'))


# --- Маршрути для звітів ---
@report_bp.route('/new', methods=['GET', 'POST'])
@login_required
def create_report():
    if request.method == 'POST':
        try:
            enterprise_id, report_year, report_month = request.form.get('enterprise_id',type=int), request.form.get('report_year',type=int), request.form.get('report_month',type=int)
            if MonthlyReport.query.filter_by(enterprise_id=enterprise_id, report_year=report_year, report_month=report_month).first():
                flash('Звіт для цього підприємства за вказаний місяць вже існує.', 'danger')
                return redirect(url_for('report.create_report'))
            new_report = MonthlyReport(user_id=current_user.id, enterprise_id=enterprise_id, report_year=report_year, report_month=report_month, production_total=request.form.get('production_total'), production_researched=request.form.get('production_researched'), waste_utilized_ton=request.form.get('waste_utilized_ton', 0.0, type=float), shortcomings_violations=request.form.get('shortcomings_violations'))
            db.session.add(new_report)
            for animal in ['vrh','svyni','vivci_kozy','koni','ptytsia','inshi']:
                report_data_entry = ReportData(report=new_report, animal_type=animal, received=request.form.get(f'{animal}_received',0,type=int), diseases_registered=request.form.get(f'{animal}_diseases_registered',0,type=int), disease_sybirka=request.form.get(f'{animal}_disease_sybirka',0,type=int), disease_tuberkuloz=request.form.get(f'{animal}_disease_tuberkuloz',0,type=int), disease_brutseloz=request.form.get(f'{animal}_disease_brutseloz',0,type=int), disease_lepto=request.form.get(f'{animal}_disease_lepto',0,type=int), disease_beshykha_svynei=request.form.get(f'{animal}_disease_beshykha_svynei',0,type=int), disease_chuma_svynei=request.form.get(f'{animal}_disease_chuma_svynei',0,type=int), disease_nezarazni=request.form.get(f'{animal}_disease_nezarazni',0,type=int), disease_inshi_zarazni=request.form.get(f'{animal}_disease_inshi_zarazni',0,type=int), died_from_trauma=request.form.get(f'{animal}_died_from_trauma',0,type=int), sent_to_sanbiynia=request.form.get(f'{animal}_sent_to_sanbiynia',0,type=int), expert_sybirka=request.form.get(f'{animal}_expert_sybirka',0,type=int), expert_tuberkuloz_total=request.form.get(f'{animal}_expert_tuberkuloz_total',0,type=int), expert_tuberkuloz_util=request.form.get(f'{animal}_expert_tuberkuloz_util',0,type=int), expert_tuberkuloz_sanpererobka=request.form.get(f'{animal}_expert_tuberkuloz_sanpererobka',0,type=int), expert_leikoz=request.form.get(f'{animal}_expert_leikoz',0,type=int), expert_tsystytserkoz_finoz=request.form.get(f'{animal}_expert_tsystytserkoz_finoz',0,type=int), expert_ekhinokokoz=request.form.get(f'{animal}_expert_ekhinokokoz',0,type=int), expert_fastsioloz=request.form.get(f'{animal}_expert_fastsioloz',0,type=int), expert_brutseloz=request.form.get(f'{animal}_expert_brutseloz',0,type=int), expert_trykhineloz=request.form.get(f'{animal}_expert_trykhineloz',0,type=int), expert_leptospiroz=request.form.get(f'{animal}_expert_leptospiroz',0,type=int), expert_inshi_zarazni=request.form.get(f'{animal}_expert_inshi_zarazni',0,type=int), expert_inshi_invaziyni=request.form.get(f'{animal}_expert_inshi_invaziyni',0,type=int), expert_nezarazni_util=request.form.get(f'{animal}_expert_nezarazni_util',0,type=int), expert_nezarazni_prompererobka=request.form.get(f'{animal}_expert_nezarazni_prompererobka',0,type=int), expert_nezarazni_na_utylzavod=request.form.get(f'{animal}_expert_nezarazni_na_utylzavod',0,type=int), expert_zneshkodzhennia_tush=request.form.get(f'{animal}_expert_zneshkodzhennia_tush',0,type=int), expert_utylzavod_holiv=request.form.get(f'{animal}_expert_utylzavod_holiv',0,type=int))
                db.session.add(report_data_entry)
            db.session.commit()
            flash('Звіт успішно створено!', 'success')
            return redirect(url_for('report.archive'))
        except Exception as e:
            db.session.rollback(); flash(f'Під час створення звіту сталася помилка: {e}', 'danger'); traceback.print_exc()
    user_enterprises = Enterprise.query.filter_by(user_id=current_user.id).all()
    return render_template('create_report.html', enterprises=user_enterprises, current_year=datetime.datetime.now().year, uk_months=uk_months)

@report_bp.route('/archive')
@login_required
def archive():
    reports = MonthlyReport.query.filter_by(author=current_user).order_by(MonthlyReport.report_year.desc(), MonthlyReport.report_month.desc()).all()
    return render_template('report_archive.html', reports=reports, uk_months=uk_months)

@report_bp.route('/view/<int:report_id>')
@login_required
def view_report(report_id):
    report = db.session.get(MonthlyReport, report_id)
    if not report or (report.user_id != current_user.id and not current_user.is_manager):
        flash('Звіт не знайдено або у вас немає доступу.', 'danger')
        return redirect(url_for('report.archive'))
    report_data_dict = {data.animal_type: data for data in report.report_data}
    animal_types_display = {'vrh':'ВРХ', 'svyni':'Свині', 'vivci_kozy':'Вівці і кози', 'koni':'Коні', 'ptytsia':'Птиця', 'inshi':'Інші види'}
    return render_template('view_report.html', report=report, report_data=report_data_dict, uk_months=uk_months, animal_types_display=animal_types_display)

@report_bp.route('/download_excel/<int:report_id>')
@login_required
def download_excel_report(report_id):
    report = db.session.get(MonthlyReport, report_id)
    if not report or (report.user_id != current_user.id and not current_user.is_manager):
        flash('Звіт не знайдено або у вас немає доступу.', 'danger')
        return redirect(url_for('report.archive'))
    
    workbook = openpyxl.Workbook()
    
    virtual_workbook = io.BytesIO(); workbook.save(virtual_workbook); virtual_workbook.seek(0)
    filename = f'Zvit_{report.id}_{report.enterprise.name}_{report.report_year}_{report.report_month}.xlsx'
    return send_file(virtual_workbook, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@document_bp.route('/')
@login_required
def list_docs():
    documents = Document.query.order_by(Document.upload_date.desc()).all()
    return render_template('document_list.html', documents=documents)

@document_bp.route('/upload', methods=['GET', 'POST'])
@login_required
@manager_required
def upload_doc():
    if request.method == 'POST':
        file, title, description = request.files.get('file'), request.form.get('title'), request.form.get('description')
        if not all([file, file.filename, title]):
            flash('Будь ласка, заповніть назву та виберіть файл.', 'danger')
            return redirect(request.url)
        if allowed_file(file.filename, ALLOWED_DOC_EXTENSIONS):
            try:
                filename = secure_filename(file.filename)
                save_path = os.path.join(UPLOAD_FOLDER_DOCS, filename)
                file.save(save_path)
                relative_path = os.path.join('uploads', 'documents', filename).replace('\\', '/')
                new_doc = Document(title=title, description=description, filename=filename, filepath=relative_path, user_id=current_user.id)
                db.session.add(new_doc)
                db.session.commit()
                flash(f'Документ "{title}" успішно завантажено!', 'success')
                return redirect(url_for('document.list_docs'))
            except Exception as e:
                flash(f'Під час завантаження файлу сталася помилка: {e}', 'danger'); traceback.print_exc()
                return redirect(request.url)
        else:
            flash('Недопустимий тип файлу. Дозволено: doc, docx, xls, xlsx, pdf, txt.', 'danger')
            return redirect(request.url)
    return render_template('upload_document.html')

@document_bp.route('/download/<int:doc_id>')
@login_required
def download_doc(doc_id):
    doc = db.session.get(Document, doc_id)
    if not doc:
        flash('Документ не знайдено.', 'danger')
        return redirect(url_for('document.list_docs'))
    try:
        full_path = os.path.join(basedir, 'static', doc.filepath)
        return send_file(full_path, as_attachment=True)
    except Exception as e:
        flash(f'Не вдалося завантажити файл: {e}', 'danger'); traceback.print_exc()
        return redirect(url_for('document.list_docs'))

@document_bp.route('/delete/<int:doc_id>', methods=['POST'])
@login_required
@manager_required
def delete_doc(doc_id):
    if not current_user.is_admin:
        flash('Тільки адміністратор може видаляти документи.', 'danger')
        return redirect(url_for('document.list_docs'))
    
    doc = db.session.get(Document, doc_id)
    if doc:
        try:
            local_path = os.path.join(app.static_folder, doc.filepath)
            if os.path.exists(local_path):
                os.remove(local_path)
            db.session.delete(doc)
            db.session.commit()
            flash(f'Документ "{doc.title}" успішно видалено.', 'success')
        except Exception as e:
            db.session.rollback(); flash(f'Помилка видалення документа: {e}', 'danger'); traceback.print_exc()
    else:
        flash('Документ не знайдено.', 'danger')
    return redirect(url_for('document.list_docs'))


# --- Реєстрація всіх Blueprints ---
app.register_blueprint(main_bp)
app.register_blueprint(auth_bp)
app.register_blueprint(photo_bp)
app.register_blueprint(admin_bp)
app.register_blueprint(report_bp)
app.register_blueprint(document_bp)
app.register_blueprint(safety_bp)


# ====================================================================
# 7. Створення бази даних та початкових користувачів
# ====================================================================
@app.cli.command("init-db")
def init_db_command():
    """Створює/очищує базу даних та створює початкових користувачів."""
    db.drop_all()
    db.create_all()
    if not User.query.filter_by(username='admin').first():
        admin_user = User(username='admin', role='admin', is_active=True)
        admin_user.set_password('adminpassword')
        db.session.add(admin_user)
        print("Створено адміністратора.")
    if not User.query.filter_by(username='manager').first():
        manager_user = User(username='manager', role='manager', is_active=True)
        manager_user.set_password('managerpass')
        db.session.add(manager_user)
        print("Створено керівника.")
    db.session.commit()
    print("Ініціалізовано базу даних.")


if __name__ == '__main__':
    app.run(debug=True)